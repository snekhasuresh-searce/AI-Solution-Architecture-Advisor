import io

from docx import Document
from pypdf import PdfReader

from advisor.export import inline_runs, parse, to_docx, to_pdf

SAMPLE = """# Company Website

**Status:** Approved by reviewer
**Run:** abc12345

## 2. Solution design

| Category | Choice | From |
| --- | --- | --- |
| frontend_framework | Astro | frontend, uiux |
| state_data_fetching | Zustand → none | frontend |

- **frontend**: needed for frontend
- *uiux*: needed

> Human guidance recommended: low confidence.

---
_Saved to `outputs/x.md`._
"""


def test_parse_blocks():
    kinds = [b.kind for b in parse(SAMPLE)]
    assert kinds == ["heading", "para", "heading", "table", "bullets", "quote", "rule", "para"]
    table = parse(SAMPLE)[3]
    assert table.rows[0] == ["Category", "Choice", "From"] and len(table.rows) == 3


def test_inline_keeps_snake_case():
    assert inline_runs("state_data_fetching") == [("", "state_data_fetching")]
    assert inline_runs("**a** and _b_ and `c`") == [("b", "a"), ("", " and "), ("i", "b"), ("", " and "), ("c", "c")]


def test_docx_contains_content():
    doc = Document(io.BytesIO(to_docx(SAMPLE)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Company Website" in text and "Approved by reviewer" in text
    assert doc.tables[0].cell(1, 1).text == "Astro"
    assert doc.core_properties.title == "Company Website"


def test_pdf_contains_content():
    reader = PdfReader(io.BytesIO(to_pdf(SAMPLE)))
    text = "".join(page.extract_text() for page in reader.pages)
    assert "Company Website" in text and "Astro" in text and "state_data_fetching" in text

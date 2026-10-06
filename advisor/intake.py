"""Read a requirement from .txt/.md/.pdf/.docx into plain text."""

from __future__ import annotations

from pathlib import Path


def read_requirement_file(path: str) -> str:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in {".txt", ".md"}:
        return p.read_text(encoding="utf-8")
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(str(p)).pages)
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(p))
        parts = [para.text for para in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(t for t in parts if t.strip())
    raise ValueError(f"Unsupported file type: {suffix} (use .txt, .md, .pdf or .docx)")

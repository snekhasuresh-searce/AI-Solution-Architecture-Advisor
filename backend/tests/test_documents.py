"""Editable recommendations: save, conflicts, validation, version history, access rules and exports."""

import io
import os

os.environ["ADVISOR_PROVIDER"] = "mock"

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

RUN = "abcdef12"
REPORT = """# Company website

**Status:** Approved by reviewer
**Run:** abcdef12

## 1. Executive summary

A *responsive* site with **three** pages.

| Page | Purpose |
|---|---|
| Home | Landing |
"""
PNG = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def text(value, *marks):
    return {"type": "text", "text": value, **({"marks": list(marks)} if marks else {})}


def para(*children, **attrs):
    return {"type": "paragraph", **({"attrs": attrs} if attrs else {}), "content": list(children)}


def cell(kind, value, **attrs):
    return {"type": kind, "attrs": {"colspan": 1, "rowspan": 1, **attrs}, "content": [para(text(value))]}


RICH = {"type": "doc", "content": [
    {"type": "heading", "attrs": {"level": 1, "textAlign": "center"}, "content": [text("Project Overview")]},
    para(text("Bold", {"type": "bold"}), text(" and "), text("red", {"type": "textStyle", "attrs": {"color": "#ff0000",
         "fontSize": "14pt", "fontFamily": "Georgia"}}), text(" and "),
         text("marked", {"type": "highlight", "attrs": {"color": "#fff59d"}}), text(" "),
         text("under", {"type": "underline"}), text(" "), text("gone", {"type": "strike"}), text(" "),
         text("a link", {"type": "link", "attrs": {"href": "https://example.com"}}),
         {"type": "hardBreak"}, text("next line"), lineHeight="1.5", textAlign="justify"),
    {"type": "heading", "attrs": {"level": 2}, "content": [text("Functional Requirements")]},
    {"type": "orderedList", "attrs": {"start": 1}, "content": [
        {"type": "listItem", "content": [para(text("First")), {"type": "bulletList", "content": [
            {"type": "listItem", "content": [para(text("Nested bullet"))]}]}]},
        {"type": "listItem", "content": [para(text("Second"))]},
    ]},
    {"type": "orderedList", "attrs": {"start": 1}, "content": [{"type": "listItem", "content": [para(text("Restart"))]}]},
    {"type": "table", "content": [
        {"type": "tableRow", "content": [cell("tableHeader", "Role"), cell("tableHeader", "Can")]},
        {"type": "tableRow", "content": [cell("tableCell", "Admin", rowspan=2), cell("tableCell", "Everything")]},
        {"type": "tableRow", "content": [cell("tableCell", "Approve")]},
        {"type": "tableRow", "content": [cell("tableCell", "Merged across", colspan=2)]},
    ]},
    {"type": "horizontalRule"},
    {"type": "image", "attrs": {"src": PNG, "alt": "pixel"}},
    {"type": "image", "attrs": {"src": f"/api/runs/{RUN}/diagram.svg", "alt": "Architecture"}},
    {"type": "blockquote", "content": [para(text("A quote"))]},
]}


@pytest.fixture
def env(tmp_path, monkeypatch):
    from dataclasses import replace

    from advisor import architecture, report, storage

    settings = replace(storage.settings, database_url="", output_dir=str(tmp_path),
                       db_path=str(tmp_path / "runs.sqlite"))
    monkeypatch.setattr(storage, "settings", settings)
    monkeypatch.setattr(report, "settings", settings)
    monkeypatch.setattr(architecture, "settings", settings)
    return settings


def login(email, role=None):
    from advisor.api import app

    client = TestClient(app, headers={"X-Advisor-Client": "test"})
    res = client.post("/api/auth/dev-login", json={"email": email, **({"role": role} if role else {})})
    assert res.status_code == 200, res.text
    return client, res.json()["id"]


@pytest.fixture
def owner(env):
    """A consultant who owns run RUN, which has only its AI-generated report."""
    from sqlalchemy import delete

    from advisor import db, documents, report, storage

    client, user_id = login("writer@searce.com")
    report.save(RUN, REPORT)
    storage.log_run(run_id=RUN, session_id="s", requirement="r", domains=["frontend"], agents=["frontend"],
                    status="APPROVED", overall=90.0, rounds=1, payload={}, owner_id=user_id)
    documents.ensure_tables()
    with db.engine().begin() as conn:  # the app database outlives tmp_path between tests
        conn.execute(delete(documents.documents).where(documents.documents.c.run_id == RUN))
        conn.execute(delete(documents.document_versions).where(documents.document_versions.c.run_id == RUN))
    return client


def save(client, content=RICH, base=None, title="Edited website", manual=False):
    return client.put(f"/api/runs/{RUN}/document",
                      json={"title": title, "content": content, "base_version": base, "manual": manual})


def test_unedited_run_has_no_document_and_exports_the_original(owner):
    state = owner.get(f"/api/runs/{RUN}/document").json()
    assert state == {"run_id": RUN, "editable": True, "document": None}
    md = owner.get(f"/api/runs/{RUN}/export/md")
    assert md.status_code == 200 and md.text == REPORT
    assert owner.get(f"/api/runs/{RUN}/export/docx").content.startswith(b"PK")


def test_save_versions_and_conflicts(owner):
    first = save(owner)
    assert first.status_code == 200, first.text
    doc = first.json()
    assert doc["version"] >= 1 and doc["title"] == "Edited website" and doc["updated_by"] == "writer@searce.com"
    assert owner.get(f"/api/runs/{RUN}/document").json()["document"]["content"] == RICH

    # A second "first save" (another tab that loaded the original) is a conflict, and so is a stale version.
    assert save(owner).status_code == 409
    second = save(owner, base=doc["version"], title="Renamed", manual=True)
    assert second.status_code == 200 and second.json()["version"] == doc["version"] + 1
    assert save(owner, base=doc["version"]).status_code == 409

    # The run history shows the new title.
    run = next(r for r in owner.get("/api/runs").json() if r["run_id"] == RUN)
    assert run["title"] == "Renamed" and run["edited"] is True

    versions = owner.get(f"/api/runs/{RUN}/document/versions").json()
    assert [v["reason"] for v in versions][:2] == ["manual", "auto"]
    snap = owner.get(f"/api/runs/{RUN}/document/versions/{versions[1]['id']}").json()
    assert snap["title"] == "Edited website" and snap["content"] == RICH
    assert owner.get(f"/api/runs/{RUN}/document/versions/nope").status_code == 404

    # Reset: back to the AI report, the edit kept in the history; saving again continues the numbering.
    assert owner.delete(f"/api/runs/{RUN}/document").json() == {"reset": True}
    assert owner.get(f"/api/runs/{RUN}/document").json()["document"] is None
    assert owner.get(f"/api/runs/{RUN}/document/versions").json()[0]["reason"] == "reset"
    again = save(owner)
    assert again.status_code == 200 and again.json()["version"] == second.json()["version"] + 1


@pytest.mark.parametrize("content, why", [
    ({"type": "doc"}, "no content list"),
    ({"type": "doc", "content": [{"type": "script"}]}, "unknown node"),
    ({"type": "doc", "content": [para(text("x", {"type": "link", "attrs": {"href": "javascript:alert(1)"}}))]}, "js link"),
    ({"type": "doc", "content": [{"type": "image", "attrs": {"src": "file:///etc/passwd"}}]}, "file image"),
    ({"type": "doc", "content": [para(text("x", {"type": "comment"}))]}, "unknown mark"),
])
def test_invalid_content_is_rejected(owner, content, why):
    res = save(owner, content=content)
    assert res.status_code == 422, why


def test_blank_title_is_rejected(owner):
    assert save(owner, title="   ").status_code == 422


def test_access_rules(owner):
    other, _ = login("someone@searce.com")
    assert other.get(f"/api/runs/{RUN}/document").status_code == 404  # cannot even see the run
    assert save(other).status_code == 404

    reviewer, _ = login("rev@searce.com", role="reviewer")
    state = reviewer.get(f"/api/runs/{RUN}/document")
    assert state.status_code == 200 and state.json()["editable"] is False
    assert save(reviewer).status_code == 403
    assert reviewer.delete(f"/api/runs/{RUN}/document").status_code == 403

    admin, _ = login("boss@searce.com")
    assert admin.get(f"/api/runs/{RUN}/document").json()["editable"] is True
    assert save(admin).status_code == 200


def test_writes_need_the_client_header(owner):
    from advisor.api import app

    bare = TestClient(app)
    bare.cookies = owner.cookies
    assert bare.put(f"/api/runs/{RUN}/document", json={"title": "x", "content": RICH}).status_code == 403


def test_exports_keep_the_edits(owner):
    assert save(owner).status_code == 200

    md = owner.get(f"/api/runs/{RUN}/export/md")
    assert md.status_code == 200 and md.headers["content-type"].startswith("text/markdown")
    for expected in ("# Edited website", "## Project Overview", "### Functional Requirements", "**Bold**",
                     "<u>under</u>", "~~gone~~", "[a link](https://example.com)", "1. First", "   - Nested bullet",
                     "| Role | Can |", "| Admin | Everything |", "---", "![pixel](data:image/png", "> A quote"):
        assert expected in md.text, expected

    pdf = owner.get(f"/api/runs/{RUN}/export/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")

    res = owner.get(f"/api/runs/{RUN}/export/docx")
    assert res.status_code == 200 and f"recommendation_{RUN}.docx" in res.headers["content-disposition"]
    from docx import Document

    doc = Document(io.BytesIO(res.content))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "Edited website" in body and "Project Overview" in body and "Nested bullet" in body
    styles = {p.style.name for p in doc.paragraphs}
    assert {"Title", "Heading 1", "Heading 2", "List Number", "List Bullet 2", "Intense Quote"} <= styles
    red = next(r for p in doc.paragraphs for r in p.runs if r.text == "red")
    assert str(red.font.color.rgb) == "FF0000" and red.font.size.pt == 14 and red.font.name == "Georgia"
    assert any(r.font.strike for p in doc.paragraphs for r in p.runs if r.text == "gone")
    assert "https://example.com" in [rel.target_ref for rel in doc.part.rels.values() if rel.is_external]
    # Two numbered lists, each starting at 1: separate numbering instances.
    nums = {p._p.pPr.numPr.numId.val for p in doc.paragraphs if p.style.name == "List Number"}
    assert len(nums) == 2
    table = doc.tables[0]
    assert table.cell(1, 0).text == table.cell(2, 0).text == "Admin"           # row span
    assert table.cell(3, 0).text == table.cell(3, 1).text == "Merged across"   # column span
    assert len(doc.inline_shapes) == 1  # the pasted image (body only; the logos are in the header)
    assert "[Image: Architecture]" in body  # this run has no architecture JSON to draw the diagram from


def test_doc_export_helpers():
    from advisor.doc_export import hex_color, points, table_cells

    assert hex_color("#abc") == "AABBCC" and hex_color("rgb(255, 0, 16)") == "FF0010" and hex_color("red") is None
    assert points("16px") == 12 and points("11pt") == 11 and points("1000pt") == 96 and points("x") is None
    cells, rows, cols, header_rows = table_cells(RICH["content"][5])
    assert (rows, cols, header_rows) == (4, 2, 1)
    assert [(c.row, c.col) for c in cells] == [(0, 0), (0, 1), (1, 0), (1, 1), (2, 1), (3, 0)]

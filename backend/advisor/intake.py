"""Read a requirement from .txt/.md/.pdf/.docx/.pptx into plain text."""

from __future__ import annotations

from pathlib import Path


def read_requirement_file(path: str) -> str:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in {".txt", ".md"}:
        return p.read_text(encoding="utf-8")
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(str(p)).pages)
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(p))
        parts = [para.text for para in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(t for t in parts if t.strip())
    if suffix == ".pptx":
        return _read_pptx(p)
    raise ValueError(f"Unsupported file type: {suffix} (use .txt, .md, .pdf, .docx or .pptx)")


def _read_pptx(p: Path) -> str:
    """Slide by slide: text boxes, tables and grouped shapes in order, then the speaker notes."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def shape_text(shape) -> list[str]:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            return [t for child in shape.shapes for t in shape_text(child)]
        if shape.has_table:
            return [" | ".join(cell.text for cell in row.cells) for row in shape.table.rows]
        if shape.has_text_frame:
            return [para.text for para in shape.text_frame.paragraphs]
        return []

    parts: list[str] = []
    for number, slide in enumerate(Presentation(str(p)).slides, start=1):
        lines = [t for shape in slide.shapes for t in shape_text(shape) if t.strip()]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                lines.append(f"Notes: {notes}")
        if lines:
            parts.append(f"Slide {number}\n" + "\n".join(lines))
    return "\n\n".join(parts)

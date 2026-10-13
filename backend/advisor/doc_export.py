"""Export an edited recommendation (editor JSON from documents.py) to DOCX, PDF and Markdown.

Same branded page frame and table look as export.py, which renders the AI-generated Markdown; this one
also keeps what the editor adds: underline, strikethrough, fonts, sizes, colours, highlights, alignment,
line spacing, numbered and nested lists, merged table cells, links and pasted images.

Heading levels: the document title is the Title, editor Heading 1 is a numbered report section
(Markdown ##), Heading 2 a subsection (###) and Heading 3 (####).
"""

from __future__ import annotations

import base64
import binascii
import io
import re
from dataclasses import dataclass, replace

from .export import (ACCENT, BAND, NAVY, RULE, _ASCII, _col_widths, _fonts, docx_border, docx_shade,
                     docx_style_table, new_docx, pdf_frame)

DIAGRAM = re.compile(r"^/api/runs/([0-9a-f]{8})/diagram\.(svg|png)$")
DATA_URI = re.compile(r"^data:image/(png|jpe?g|gif);base64,(.+)$", re.I | re.S)


# ------------------------------------------------------------------ inline
@dataclass(frozen=True)
class Run:
    text: str                 # "\n" = line break
    bold: bool = False
    italic: bool = False
    underline: bool = False
    strike: bool = False
    code: bool = False
    link: str | None = None
    color: str | None = None  # RRGGBB
    highlight: str | None = None
    font: str | None = None
    size: float | None = None  # points


def hex_color(value) -> str | None:
    """'#abc', '#aabbcc' or 'rgb(r, g, b)' -> 'AABBCC'."""
    s = str(value or "").strip().lower()
    if m := re.fullmatch(r"#([0-9a-f]{3})", s):
        return "".join(c * 2 for c in m.group(1)).upper()
    if m := re.fullmatch(r"#([0-9a-f]{6})([0-9a-f]{2})?", s):
        return m.group(1).upper()
    if m := re.fullmatch(r"rgba?\((\d+),\s*(\d+),\s*(\d+)(?:,\s*[\d.]+)?\)", s):
        return "".join(f"{min(int(v), 255):02X}" for v in m.groups())
    return None


def points(value) -> float | None:
    """'12pt' / '16px' / '12' -> points, kept to a sane range."""
    m = re.fullmatch(r"\s*([\d.]+)\s*(pt|px)?\s*", str(value or ""))
    if not m:
        return None
    try:
        size = float(m.group(1)) * (0.75 if m.group(2) == "px" else 1)
    except ValueError:
        return None
    return min(max(size, 6), 96)


def font_name(value) -> str | None:
    """First family of a CSS font-family list."""
    name = str(value or "").split(",")[0].strip().strip("'\"")
    return name or None


def runs_of(node: dict) -> list[Run]:
    runs: list[Run] = []
    for child in node.get("content") or []:
        if child.get("type") == "hardBreak":
            runs.append(Run("\n"))
            continue
        if child.get("type") != "text" or not child.get("text"):
            continue
        run = Run(child["text"])
        for mark in child.get("marks") or []:
            attrs = mark.get("attrs") or {}
            kind = mark.get("type")
            if kind in ("bold", "italic", "underline", "strike", "code"):
                run = replace(run, **{kind: True})
            elif kind == "link":
                run = replace(run, link=str(attrs.get("href") or "") or None)
            elif kind == "highlight":
                run = replace(run, highlight=hex_color(attrs.get("color")) or "FFF59D")
            elif kind == "textStyle":
                run = replace(run, color=hex_color(attrs.get("color")) or run.color,
                              highlight=hex_color(attrs.get("backgroundColor")) or run.highlight,
                              font=font_name(attrs.get("fontFamily")) or run.font,
                              size=points(attrs.get("fontSize")) or run.size)
        runs.append(run)
    return runs


def plain(node: dict) -> str:
    """All text inside a node, for table column sizing."""
    if node.get("type") == "text":
        return node.get("text", "")
    sep = "\n" if node.get("type") not in ("paragraph", "heading") else ""
    return sep.join(plain(c) for c in node.get("content") or [])


def line_height(attrs: dict) -> float | None:
    try:
        value = float(attrs.get("lineHeight") or 0)
    except (TypeError, ValueError):
        return None
    return min(max(value, 0.8), 4) if value else None


# ------------------------------------------------------------------ tables
@dataclass
class Cell:
    row: int
    col: int
    rowspan: int
    colspan: int
    header: bool
    node: dict


def table_cells(table: dict) -> tuple[list[Cell], int, int, int]:
    """Place cells on a grid, honouring row and column spans. Returns (cells, rows, cols, header_rows)."""
    taken: set[tuple[int, int]] = set()
    cells: list[Cell] = []
    rows = [r for r in table.get("content") or [] if r.get("type") == "tableRow"]
    for r, row in enumerate(rows):
        c = 0
        for node in row.get("content") or []:
            while (r, c) in taken:
                c += 1
            attrs = node.get("attrs") or {}
            rs = max(1, min(int(attrs.get("rowspan") or 1), len(rows) - r))
            cs = max(1, int(attrs.get("colspan") or 1))
            taken.update((r + dr, c + dc) for dr in range(rs) for dc in range(cs))
            cells.append(Cell(r, c, rs, cs, node.get("type") == "tableHeader", node))
            c += cs
    ncols = max((c for _, c in taken), default=-1) + 1
    header_rows = 0
    for r in range(len(rows)):
        row_cells = [x for x in cells if x.row == r]
        if row_cells and all(x.header for x in row_cells):
            header_rows += 1
        else:
            break
    return cells, len(rows), ncols, header_rows


def column_widths(cells: list[Cell], nrows: int, ncols: int, total: float, char_w: float) -> list[float]:
    grid = [["" for _ in range(ncols)] for _ in range(nrows)]
    for cell in cells:
        if cell.colspan == 1:
            grid[cell.row][cell.col] = plain(cell.node)
    return _col_widths(grid, total, char_w=char_w) if ncols else []


def image_bytes(src: str) -> bytes | None:
    if m := DATA_URI.match(src or ""):
        try:
            return base64.b64decode(m.group(2), validate=False)
        except (binascii.Error, ValueError):
            return None
    return None


def diagram_for(src: str, run_id: str, architecture: dict | None) -> dict | None:
    """The run's own architecture diagram (never another run's)."""
    m = DIAGRAM.match(src or "")
    return architecture if m and m.group(1) == run_id and architecture else None


# -------------------------------------------------------------------- DOCX
def to_docx(content: dict, title: str, run_id: str, architecture: dict | None = None) -> bytes:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Emu, Pt, RGBColor
    from docx.text.run import Run as DocxRun

    doc = new_docx(title)
    heading = doc.add_heading(level=0)
    heading.add_run(title)
    docx_border(heading, "bottom", ACCENT)
    sec = doc.sections[-1]
    page_width = sec.page_width - sec.left_margin - sec.right_margin
    align = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
             "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT}
    numbering = doc.part.numbering_part.numbering_definitions._numbering

    def style_run(r, run: Run, *, size=None, bold=False, color=None) -> None:
        r.bold = bold or run.bold or None
        r.italic = run.italic or None
        r.underline = run.underline or bool(run.link) or None
        r.font.strike = run.strike or None
        if run.code:
            r.font.name = "Consolas"
        elif run.font:
            r.font.name = run.font
        if run.size or size:
            r.font.size = Pt(run.size or size)
        rgb = run.color or color or ("0068D6" if run.link else None)
        if rgb:
            r.font.color.rgb = RGBColor.from_string(rgb)
        if run.highlight:
            docx_shade(r, run.highlight)

    def add_runs(paragraph, runs: list[Run], **fmt) -> None:
        for run in runs:
            if run.text == "\n":
                paragraph.add_run().add_break()
            elif run.link:
                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), paragraph.part.relate_to(run.link, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                r = DocxRun(OxmlElement("w:r"), paragraph)
                r.text = run.text
                style_run(r, run, **fmt)
                link.append(r._r)
                paragraph._p.append(link)
            else:
                style_run(paragraph.add_run(run.text), run, **fmt)

    def paragraph_format(paragraph, attrs: dict) -> None:
        if attrs.get("textAlign") in align:
            paragraph.alignment = align[attrs["textAlign"]]
        if lh := line_height(attrs):
            paragraph.paragraph_format.line_spacing = lh

    def new_list(style: str) -> int:
        """A fresh numbering instance, so each numbered list starts again at 1."""
        num_id = doc.styles[style].element.pPr.numPr.numId.val
        abstract = numbering.num_having_numId(num_id).abstractNumId.val
        num = numbering.add_num(abstract)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return num.numId

    def picture(container, data: bytes, max_width: int) -> None:
        run = container.add_paragraph().add_run()
        shape = run.add_picture(io.BytesIO(data))
        if shape.width > max_width:
            shape.height = Emu(int(shape.height * max_width / shape.width))
            shape.width = Emu(max_width)

    def blocks(container, nodes: list[dict], width: int, depth: int = 0, list_style: str | None = None,
               num_id: int | None = None, in_cell: bool = False) -> None:
        for node in nodes:
            kind, attrs = node.get("type"), node.get("attrs") or {}
            if kind == "heading":
                level = min(max(int(attrs.get("level") or 1), 1), 3)
                p = container.add_paragraph(style=f"Heading {level}")
                add_runs(p, runs_of(node))
                paragraph_format(p, attrs)
                if level == 1 and not in_cell:
                    docx_border(p, "bottom", RULE)
            elif kind == "paragraph":
                p = container.add_paragraph()
                add_runs(p, runs_of(node), size=9 if in_cell else None)
                paragraph_format(p, attrs)
                if in_cell:
                    p.paragraph_format.space_after = Pt(0)
            elif kind in ("bulletList", "orderedList"):
                level = min(depth, 2)
                base = "List Bullet" if kind == "bulletList" else "List Number"
                style = base if level == 0 else f"{base} {level + 1}"
                nid = new_list(style) if kind == "orderedList" else None
                blocks(container, node.get("content") or [], width, depth + 1, style, nid, in_cell)
            elif kind == "listItem":
                children = node.get("content") or []
                first = children[0] if children and children[0].get("type") == "paragraph" else None
                if first is not None:
                    p = container.add_paragraph(style=list_style)
                    add_runs(p, runs_of(first), size=9 if in_cell else None)
                    paragraph_format(p, first.get("attrs") or {})
                    if num_id is not None:
                        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
                        num_pr.get_or_add_numId().val = num_id
                        num_pr.get_or_add_ilvl().val = 0
                    children = children[1:]
                blocks(container, children, width, depth, list_style, num_id, in_cell)
            elif kind == "blockquote":
                for child in node.get("content") or []:
                    if child.get("type") == "paragraph":
                        p = container.add_paragraph(style="Intense Quote")
                        add_runs(p, runs_of(child))
                    else:
                        blocks(container, [child], width, depth, list_style, num_id, in_cell)
            elif kind == "codeBlock":
                p = container.add_paragraph()
                r = p.add_run(plain(node))
                r.font.name, r.font.size = "Consolas", Pt(9)
                docx_shade(r, "F6F8FA")
            elif kind == "horizontalRule":
                p = container.add_paragraph()
                docx_border(p, "bottom", RULE)
            elif kind == "image":
                src = str(attrs.get("src") or "")
                if arch := diagram_for(src, run_id, architecture):
                    from .diagram import to_png

                    picture(container, to_png(arch), width)
                elif data := image_bytes(src):
                    try:
                        picture(container, data, width)
                    except Exception:  # noqa: BLE001 - an image Word cannot read
                        container.add_paragraph(f"[Image: {attrs.get('alt') or 'not supported in Word'}]")
                else:
                    add_runs(container.add_paragraph(), [Run(f"[Image: {attrs.get('alt') or src[:80]}]", link=src or None)])
            elif kind == "table":
                table_block(container, node, width)

    def table_block(container, node: dict, width: int) -> None:
        cells, nrows, ncols, header_rows = table_cells(node)
        if not ncols:
            return
        table = container.add_table(rows=nrows, cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        docx_style_table(table)
        widths = column_widths(cells, nrows, ncols, width / 12700, char_w=5.6)
        for r, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if r < header_rows:
                trpr.append(OxmlElement("w:tblHeader"))  # repeat header rows on each page
            for c, cell in enumerate(row.cells):
                cell.width = int(widths[c] * 12700)
                if r >= header_rows and r % 2 == 0:
                    docx_shade(cell, BAND)
        for cell in cells:
            target = table.cell(cell.row, cell.col)
            if cell.rowspan > 1 or cell.colspan > 1:
                target = target.merge(table.cell(cell.row + cell.rowspan - 1, min(cell.col + cell.colspan, ncols) - 1))
            first = target.paragraphs[0]
            cell_width = int(sum(widths[cell.col:cell.col + cell.colspan]) * 12700)
            blocks(target, cell.node.get("content") or [], cell_width, in_cell=True)
            if len(target.paragraphs) > 1 and not first.text:
                first._p.getparent().remove(first._p)  # the empty paragraph every new cell starts with
            if cell.header:
                docx_shade(target, NAVY)
                for p in target.paragraphs:
                    for r in p.runs:
                        r.bold = True
                        if r.font.color.rgb is None:
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        container.add_paragraph().paragraph_format.space_after = Pt(4)

    blocks(doc, content.get("content") or [], page_width)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------- PDF
def _pdf_font(name: str | None, regular: str) -> str | None:
    """Map an editor font to one the PDF has: Courier, Times or the report font."""
    n = (name or "").lower()
    if any(k in n for k in ("mono", "courier", "consol")):
        return "Courier"
    if any(k in n for k in ("times", "georgia", "garamond", "serif")) and "sans" not in n:
        return "Times-Roman"
    return None if not name else regular


def pdf_markup(runs: list[Run], unicode_font: bool, regular: str) -> str:
    from xml.sax.saxutils import escape, quoteattr

    out = []
    for run in runs:
        if run.text == "\n":
            out.append("<br/>")
            continue
        text = run.text
        if not unicode_font:
            for k, v in _ASCII.items():
                text = text.replace(k, v)
        s = escape(text)
        if run.code:
            s = f'<font face="Courier">{s}</font>'
        font_attrs = []
        if (face := _pdf_font(run.font, regular)) and not run.code:
            font_attrs.append(f"face={quoteattr(face)}")
        if run.size:
            font_attrs.append(f'size="{run.size:g}"')
        if run.color or run.link:
            font_attrs.append(f'color="#{run.color or "1F5FBF"}"')
        if font_attrs:
            s = f"<font {' '.join(font_attrs)}>{s}</font>"
        if run.highlight:
            s = f'<span backColor="#{run.highlight}">{s}</span>'
        for flag, tag in ((run.bold, "b"), (run.italic, "i"), (run.underline or bool(run.link), "u"),
                          (run.strike, "strike")):
            if flag:
                s = f"<{tag}>{s}</{tag}>"
        if run.link:
            s = f"<a href={quoteattr(run.link)}>{s}</a>"
        out.append(s)
    return "".join(out)


def to_pdf(content: dict, title: str, run_id: str, architecture: dict | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import (HRFlowable, Image, KeepTogether, ListFlowable, ListItem, Paragraph, Spacer,
                                    Table, TableStyle)

    regular, _bold, uni = _fonts()
    buf = io.BytesIO()
    doc, styles = pdf_frame(buf, title)
    para, quote, cell_style, head_style, line = (styles[k] for k in ("p", "q", "cell", "head", "line"))
    headings = {1: styles[2], 2: styles[3], 3: ParagraphStyle("h4", parent=styles[3], fontSize=10.5, leading=14)}
    align = {"left": TA_LEFT, "center": TA_CENTER, "right": TA_RIGHT, "justify": TA_JUSTIFY}
    page_h = A4[1] - doc.topMargin - doc.bottomMargin - 40

    def styled(base, attrs: dict):
        changes = {}
        if attrs.get("textAlign") in align:
            changes["alignment"] = align[attrs["textAlign"]]
        if lh := line_height(attrs):
            changes["leading"] = base.fontSize * 1.2 * lh
        return ParagraphStyle(f"{base.name}-x", parent=base, **changes) if changes else base

    def text(node: dict, style, attrs: dict | None = None):
        return Paragraph(pdf_markup(runs_of(node), uni, regular) or "&nbsp;", styled(style, attrs or {}))

    def image(src: str, alt: str, width: float):
        if arch := diagram_for(src, run_id, architecture):
            from .diagram import to_drawing

            return to_drawing(arch, width, page_h)
        if data := image_bytes(src):
            try:
                img = Image(io.BytesIO(data))
                scale = min(1.0, width / img.imageWidth, page_h / img.imageHeight)
                img.drawWidth, img.drawHeight = img.imageWidth * scale, img.imageHeight * scale
                img.hAlign = "LEFT"
                return img
            except Exception:  # noqa: BLE001 - an image reportlab cannot read
                pass
        label = Run(f"[Image: {alt or src[:80] or 'not available'}]", link=src if src.startswith("http") else None)
        return Paragraph(pdf_markup([label], uni, regular), quote)

    def blocks(nodes: list[dict], width: float, in_cell: bool = False) -> list:
        out = []
        base = cell_style if in_cell else para
        for node in nodes:
            kind, attrs = node.get("type"), node.get("attrs") or {}
            if kind == "heading":
                level = min(max(int(attrs.get("level") or 1), 1), 3)
                out.append(text(node, headings[level], attrs))
            elif kind == "paragraph":
                out.append(text(node, base, attrs))
            elif kind in ("bulletList", "orderedList"):
                items = [ListItem(blocks(item.get("content") or [], width - 18, in_cell), leftIndent=14)
                         for item in node.get("content") or []]
                if kind == "orderedList":
                    out.append(ListFlowable(items, bulletType="1", start=int(attrs.get("start") or 1), leftIndent=16,
                                            bulletFormat="%s.", bulletFontName=regular, bulletFontSize=base.fontSize))
                else:
                    out.append(ListFlowable(items, bulletType="bullet", start="•", leftIndent=14,
                                            bulletFontName=regular, bulletFontSize=base.fontSize))
            elif kind == "blockquote":
                for child in node.get("content") or []:
                    out += [text(child, quote)] if child.get("type") == "paragraph" else blocks([child], width, in_cell)
            elif kind == "codeBlock":
                from xml.sax.saxutils import escape

                code = escape(plain(node)).replace("\n", "<br/>").replace("  ", "&nbsp; ")
                out.append(Paragraph(f'<font face="Courier">{code}</font>', quote))
            elif kind == "horizontalRule":
                out.append(HRFlowable(width="100%", color=line, spaceBefore=6, spaceAfter=6))
            elif kind == "image":
                out.append(image(str(attrs.get("src") or ""), str(attrs.get("alt") or ""), width))
            elif kind == "table":
                out += table_block(node, width)
        return out

    def table_block(node: dict, width: float) -> list:
        cells, nrows, ncols, header_rows = table_cells(node)
        if not ncols:
            return []
        widths = column_widths(cells, nrows, ncols, width, char_w=5.0)
        grid = [["" for _ in range(ncols)] for _ in range(nrows)]
        commands = [
            ("ROWBACKGROUNDS", (0, header_rows), (-1, -1), [colors.white, colors.HexColor("#" + BAND)]),
            ("BOX", (0, 0), (-1, -1), 0.6, line),
            ("INNERGRID", (0, 0), (-1, -1), 0.4, line),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
        for cell in cells:
            last_col = min(cell.col + cell.colspan, ncols) - 1
            inner = sum(widths[cell.col:last_col + 1]) - 12
            content = cell.node.get("content") or []
            if cell.header:
                flow = [Paragraph(pdf_markup([replace(r, bold=True) for r in runs_of(c)], uni, regular) or "&nbsp;",
                                  head_style) if c.get("type") in ("paragraph", "heading") else blocks([c], inner, True)[0]
                        for c in content]
                commands.append(("BACKGROUND", (cell.col, cell.row), (last_col, cell.row + cell.rowspan - 1),
                                 colors.HexColor("#" + NAVY)))
            else:
                flow = blocks(content, inner, True)
            grid[cell.row][cell.col] = flow or ""
            if cell.rowspan > 1 or cell.colspan > 1:
                span = ((cell.col, cell.row), (last_col, cell.row + cell.rowspan - 1))
                commands.append(("SPAN", *span))
                if not cell.header and cell.rowspan > 1:  # one colour, not half of each row's band
                    commands.append(("BACKGROUND", *span, colors.white))
        table = Table(grid, colWidths=widths, repeatRows=header_rows)
        table.setStyle(TableStyle(commands))
        return [table, Spacer(1, 8)]

    story = [Paragraph(pdf_markup([Run(title)], uni, regular), styles[1]),
             HRFlowable(width="100%", color=colors.HexColor("#" + ACCENT), thickness=1.2, spaceAfter=8)]
    body = blocks(content.get("content") or [], doc.width)
    # Keep a heading with the block after it (a table or image would otherwise leave it stranded).
    i = 0
    while i < len(body):
        if isinstance(body[i], Paragraph) and body[i].style.name.startswith(("h2", "h3", "h4")) and i + 1 < len(body):
            story.append(KeepTogether(body[i:i + 2]))
            i += 2
        else:
            story.append(body[i])
            i += 1
    while story and isinstance(story[-1], Spacer):  # a trailing spacer can spill onto a blank page
        story.pop()
    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------- Markdown
def _md_escape(text: str) -> str:
    return re.sub(r"([\\`*_\[\]])", r"\\\1", text)


def md_inline(runs: list[Run], in_table: bool = False) -> str:
    # Merge neighbours with the same Markdown-visible formatting, so **a** **b** stays **ab**.
    merged: list[Run] = []
    for run in runs:
        key = (run.bold, run.italic, run.underline, run.strike, run.code, run.link)
        if merged and run.text != "\n" and merged[-1].text != "\n" and \
                (merged[-1].bold, merged[-1].italic, merged[-1].underline, merged[-1].strike, merged[-1].code,
                 merged[-1].link) == key:
            merged[-1] = replace(merged[-1], text=merged[-1].text + run.text)
        else:
            merged.append(run)
    out = []
    for run in merged:
        if run.text == "\n":
            out.append("<br>" if in_table else "  \n")
            continue
        core = run.text.strip()
        if not core:
            out.append(run.text)
            continue
        lead, trail = run.text[:len(run.text) - len(run.text.lstrip())], run.text[len(run.text.rstrip()):]
        s = f"`{core}`" if run.code else _md_escape(core)
        if run.strike:
            s = f"~~{s}~~"
        if run.italic:
            s = f"*{s}*"
        if run.bold:
            s = f"**{s}**"
        if run.underline and not run.link:
            s = f"<u>{s}</u>"
        if run.link:
            s = f"[{s}]({run.link})"
        out.append(lead + s + trail)
    text = "".join(out)
    return text.replace("|", "\\|").replace("\n", "<br>") if in_table else text


def to_markdown(content: dict, title: str, run_id: str = "", architecture: dict | None = None) -> str:
    def blocks(nodes: list[dict], indent: str = "") -> list[str]:
        out: list[str] = []
        for node in nodes:
            kind, attrs = node.get("type"), node.get("attrs") or {}
            if kind == "heading":
                level = min(max(int(attrs.get("level") or 1), 1), 3)
                out.append(f"{'#' * (level + 1)} {md_inline(runs_of(node))}")
            elif kind == "paragraph":
                line = md_inline(runs_of(node))
                line = re.sub(r"^(\s*)([#>+-]|\d+\.)(\s)", r"\1\\\2\3", line)  # not a heading or list
                out.append(indent + line.replace("\n", "\n" + indent))
            elif kind in ("bulletList", "orderedList"):
                lines = []
                for n, item in enumerate(node.get("content") or [], start=int(attrs.get("start") or 1)):
                    marker = "-" if kind == "bulletList" else f"{n}."
                    inner = blocks(item.get("content") or [], indent + " " * (len(marker) + 1))
                    first = inner[0].lstrip() if inner else ""
                    lines.append(f"{indent}{marker} {first}")
                    lines += [x for x in inner[1:]]
                out.append("\n".join(lines))
            elif kind == "blockquote":
                inner = "\n\n".join(blocks(node.get("content") or []))
                out.append("\n".join(f"> {x}" if x else ">" for x in inner.split("\n")))
            elif kind == "codeBlock":
                out.append(f"```{attrs.get('language') or ''}\n{plain(node)}\n```")
            elif kind == "horizontalRule":
                out.append("---")
            elif kind == "image":
                out.append(f"![{attrs.get('alt') or ''}]({attrs.get('src') or ''})")
            elif kind == "table":
                cells, nrows, ncols, _ = table_cells(node)
                if not ncols:
                    continue
                grid = [["" for _ in range(ncols)] for _ in range(nrows)]
                for cell in cells:
                    parts = [md_inline(runs_of(c), in_table=True) if c.get("type") in ("paragraph", "heading")
                             else plain(c).replace("\n", "<br>") for c in cell.node.get("content") or []]
                    grid[cell.row][cell.col] = "<br>".join(p for p in parts if p)
                rows = [f"| {' | '.join(r)} |" for r in grid]
                rows.insert(1, "|" + "---|" * ncols)
                out.append("\n".join(rows))
        return out

    return f"# {title}\n\n" + "\n\n".join(blocks(content.get("content") or [])) + "\n"


EXPORTERS = {
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pdf": (to_pdf, "application/pdf"),
    "md": (lambda content, title, run_id, architecture=None: to_markdown(content, title).encode("utf-8"),
           "text/markdown; charset=utf-8"),
}
